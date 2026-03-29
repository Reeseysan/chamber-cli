from __future__ import annotations

import os
from typing import TextIO

MAX_FILE_SIZE = 15 * 1024 * 1024  # 15MB
MAX_WORD_COUNT = 50_000


class DocumentError(Exception):
    """Raised for document loading issues."""
    pass


def _count_words(text: str) -> int:
    return len(text.split())


def _truncate_to_word_limit(text: str, limit: int) -> tuple[str, bool]:
    """Truncate text to word limit. Returns (text, was_truncated)."""
    words = text.split()
    if len(words) <= limit:
        return text, False
    truncated = " ".join(words[:limit])
    return truncated, True


def _extract_text(path: str) -> str:
    """Extract text from a file based on its extension."""
    ext = os.path.splitext(path)[1].lower()

    if ext in (".txt", ".md", ".csv", ".log", ".json", ".xml", ".html"):
        with open(path, "r", errors="replace") as f:
            return f.read()

    if ext == ".pdf":
        try:
            import pymupdf
        except ImportError:
            raise DocumentError(
                "PDF support requires: pip install chamber-cli[docs]"
            )
        text_parts = []
        with pymupdf.open(path) as doc:
            for page in doc:
                text_parts.append(page.get_text())
        return "\n".join(text_parts)

    if ext == ".docx":
        try:
            import docx
        except ImportError:
            raise DocumentError(
                "Word document support requires: pip install chamber-cli[docs]"
            )
        doc = docx.Document(path)
        return "\n".join(p.text for p in doc.paragraphs)

    if ext == ".xlsx":
        try:
            import openpyxl
        except ImportError:
            raise DocumentError(
                "Excel support requires: pip install chamber-cli[docs]"
            )
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        lines = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            lines.append(f"[Sheet: {sheet}]")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                lines.append(",".join(cells))
        wb.close()
        return "\n".join(lines)

    raise DocumentError(f"Unsupported file format: {ext}")


def load_document(path: str) -> str:
    """Load a document, extract text, enforce size limits. Returns formatted context string."""
    if not os.path.exists(path):
        raise DocumentError(f"File not found: {path}")

    file_size = os.path.getsize(path)
    if file_size > MAX_FILE_SIZE:
        raise DocumentError(
            f"File exceeds {MAX_FILE_SIZE // (1024*1024)}MB limit: "
            f"{file_size // (1024*1024)}MB"
        )

    text = _extract_text(path)
    filename = os.path.basename(path)
    word_count = _count_words(text)

    text, was_truncated = _truncate_to_word_limit(text, MAX_WORD_COUNT)
    truncation_note = ""
    if was_truncated:
        truncation_note = f"\n[Document truncated to first {MAX_WORD_COUNT:,} words (original: {word_count:,})]"

    return f"[DOCUMENT: {filename}]{truncation_note}\n\n{text}"


def load_from_stdin(stream: TextIO) -> str:
    """Load document content from stdin."""
    text = stream.read()
    if not text.strip():
        raise DocumentError("No content received from stdin.")

    word_count = _count_words(text)
    text, was_truncated = _truncate_to_word_limit(text, MAX_WORD_COUNT)
    truncation_note = ""
    if was_truncated:
        truncation_note = f"\n[Document truncated to first {MAX_WORD_COUNT:,} words (original: {word_count:,})]"

    return f"[DOCUMENT: stdin]{truncation_note}\n\n{text}"
